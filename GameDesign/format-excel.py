# 엑셀 서식 정리 — GameDesign/Excel/*.xlsx 전체를 서식 기준에 맞춘다
#
# 기준은 excel-table-creator 스킬의 "서식" 절 (테이블 시트 = Character.xlsx · Enum 시트 = Enum.xlsx).
# 값은 건드리지 않고 정렬·글꼴·칠·테두리·열 너비만 바꾼다 — 파이프라인은 값만 읽으므로 생성물은 그대로다.
#
# 사용: python GameDesign/format-excel.py            (전체)
#       python GameDesign/format-excel.py Item.xlsx  (지정 파일만)
# ⚠️ 대상 파일을 Excel에서 닫고 실행한다. 열려 있으면 저장이 실패한다.

import copy
import glob
import os
import sys
import unicodedata

import openpyxl
from openpyxl.styles import Alignment, Border, Color, PatternFill, Side

EXCEL_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Excel")

MARKERS = {"//", "C&S", "Type", "Min", "Max", "Default(Null)", "Ref"}
NUMBER_TYPES = {"int", "long", "float", "double", "byte", "short", "ID"}

THIN = Side(style="thin", color=Color(indexed=64))
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NO_FILL = PatternFill(fill_type=None)

# 테이블 시트: 설명 행·컬럼명 행 = 주황(테마 accent6). 글자는 전부 검정(테마 dk1)
ORANGE = PatternFill("solid", fgColor=Color(theme=9))
BLACK = Color(theme=1)

# Enum 시트 머리 세 줄: 검정 / 회색 / 옅은 회색
ENUM_HEAD = [
    (PatternFill("solid", fgColor=Color(theme=1)), Color(theme=0), True),
    (PatternFill("solid", fgColor=Color(theme=1, tint=0.35)), Color(theme=0), False),
    (PatternFill("solid", fgColor=Color(theme=2, tint=-0.1)), Color(theme=1), False),
]

MIN_WIDTH = 6
MAX_WIDTH = 60


def display_width(value):
    """한글·전각 문자는 2칸으로 센다."""
    if value is None:
        return 0
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in str(value))


def style(cell, *, horizontal=None, bold=None, color=None, fill=None, border=None):
    # 글꼴은 이름·크기를 살리고 굵기·색만 바꾼다.
    font = copy.copy(cell.font)
    if bold is not None:
        font.b = bold
    if color is not None:
        font.color = color
    cell.font = font
    cell.alignment = Alignment(horizontal=horizontal, vertical="center")
    if fill is not None:
        cell.fill = fill
    if border is not None:
        cell.border = border


def fit_widths(ws):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            w = display_width(cell.value)
            if w > widths.get(cell.column_letter, 0):
                widths[cell.column_letter] = w
    for letter, w in widths.items():
        # 굵은 글씨·여백을 위해 2칸 더한다.
        ws.column_dimensions[letter].width = max(MIN_WIDTH, min(MAX_WIDTH, w + 2))


def format_table(ws):
    markers = {}
    header_row = None
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a in MARKERS:
            markers[a] = r
        elif markers and a is None and ws.cell(r, 2).value is not None:
            header_row = r
            break
    if header_row is None or "Type" not in markers:
        print(f"  건너뜀 {ws.title}: 마커 행·컬럼명 행을 찾지 못했다")
        return

    last_col = max(c for c in range(2, ws.max_column + 1) if ws.cell(header_row, c).value is not None)

    for name, r in markers.items():
        style(ws.cell(r, 1), horizontal="center", bold=False, color=BLACK)
        for c in range(2, last_col + 1):
            if name == "//":
                style(ws.cell(r, c), horizontal="center", bold=True, color=BLACK, fill=ORANGE, border=BOX)
            else:
                style(ws.cell(r, c), horizontal="center", bold=False, color=BLACK, fill=NO_FILL, border=BOX)

    for c in range(2, last_col + 1):
        style(ws.cell(header_row, c), horizontal="left", bold=True, color=BLACK, fill=ORANGE, border=BOX)

    for c in range(2, last_col + 1):
        raw_type = str(ws.cell(markers["Type"], c).value or "")
        horizontal = "right" if raw_type in NUMBER_TYPES else "left"
        for r in range(header_row + 1, ws.max_row + 1):
            cell = ws.cell(r, c)
            if cell.value is None:
                continue
            style(cell, horizontal=horizontal, bold=False, color=BLACK)

    fit_widths(ws)


def format_enum(ws):
    for i, (fill, color, bold) in enumerate(ENUM_HEAD, start=1):
        for c in range(1, ws.max_column + 1):
            style(ws.cell(i, c), horizontal="center", bold=bold, color=color, fill=fill)
    for r in range(len(ENUM_HEAD) + 1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            if cell.value is None:
                continue
            style(cell, horizontal="center" if c == 1 else "left", bold=False, color=BLACK)
    fit_widths(ws)


def main():
    names = sys.argv[1:]
    paths = [os.path.join(EXCEL_DIR, n) for n in names] if names else sorted(glob.glob(os.path.join(EXCEL_DIR, "*.xlsx")))
    for path in paths:
        if os.path.basename(path).startswith("~$"):
            continue
        wb = openpyxl.load_workbook(path)
        is_enum = os.path.basename(path) == "Enum.xlsx"
        print(os.path.basename(path))
        for ws in wb.worksheets:
            format_enum(ws) if is_enum else format_table(ws)
        wb.save(path)


if __name__ == "__main__":
    main()
